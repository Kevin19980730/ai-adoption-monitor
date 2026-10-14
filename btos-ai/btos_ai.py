"""Census BTOS - AI use metrics by national / sector / subsector / employment size.

Downloads the Business Trends and Outlook Survey files from
https://www.census.gov/hfp/btos/data_downloads, keeps only the AI questions and writes
    output/YYYYMMDD_AI adoption metrics (BTOS + GenAI tracker).xlsx   (date = Census publication date of the latest period)
    data/history.csv                       (long format, every answer, estimate + SE)

Two AI series, NOT spliced (the question wording changed, the level jumped ~2x):
    v1  Sep-2023 .. Oct-2025 (smpdt 202319-202520): "...use AI in producing goods or services?"
        -> historical file "AI Core Questions.xlsx"
    v2  Dec-2025 ..          (smpdt 202524+):      "...use AI in any of its business functions?"
        -> current core files (National / Sector / Subsector / Employment Size Class / Sector x Size)
Q7 = AI use in the last two weeks, Q24 = expected AI use in the next six months.

BTOS is fortnightly (published every other Thursday). Run weekly: the script exits quietly
when Census has not published a new period. Use --force to rebuild anyway.
"""
import argparse
import datetime as dt
import json
import os
import sys
import urllib.request
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import genai_tracker as gtmod

HERE = Path(__file__).resolve().parent
RAW = HERE / "raw"
DATA = HERE / "data"
OUT = HERE / "output"
OUTNAME = "AI adoption metrics (BTOS + GenAI tracker)"
STATE = DATA / "state.json"
BASE = "https://www.census.gov/hfp/btos/downloads/"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/129.0 Safari/537.36"}

CURRENT = {  # level -> (file, dimension columns)
    "National": ("National.xlsx", []),
    "Sector": ("Sector.xlsx", ["Sector"]),
    "Subsector": ("Subsector.xlsx", ["Subsector"]),
    "Employment size": ("Employment Size Class.xlsx", ["Empsize"]),
    "Sector x Employment size": ("Sector by Employment Size Class.xlsx", ["Sector", "Empsize"]),
}
HIST_FILE = "AI Core Questions.xlsx"
HIST_SHEETS = {  # level -> (estimates sheet, SE sheet, dims)
    "National": ("National Estimates", "National SE", []),
    "Sector": ("Sector Estimates", "Sector SE", ["Sector"]),
    "Subsector": ("Subsector Estimates", "Subsector SE", ["Subsector"]),
    "Employment size": ("Employment Size Estimates", "Employment Size SE", ["Empsize"]),
    "Sector x Employment size": ("Sector x Employment Estimates", "Sector x Employment SE",
                                 ["Sector", "Empsize"]),
}
SUPP_FILE = "AI_Supplement_Table_2026.xlsx"
SUPP_SHEETS = {"National": ("National Response Estimates", "National Standard Errors", None),
               "Sector": ("Sector Response Estimates", "Sector Standard Errors", "Sector"),
               "Employment size": ("Employment Response Estimates", "Employment Standard Errors",
                                   "Empsize")}
AI_QIDS = {7: "AI use, last 2 weeks", 24: "AI use expected, next 6 months"}
V1 = "v1 producing goods/services"
V2 = "v2 any business function"
LEVELS = list(CURRENT)

SIZE = {"A": "1-4 employees", "B": "5-9 employees", "C": "10-19 employees",
        "D": "20-49 employees", "E": "50-99 employees", "F": "100-249 employees",
        "G": "250+ employees"}
SECTOR = {
    "11": "Agriculture, Forestry, Fishing and Hunting", "21": "Mining, Quarrying, and Oil and Gas Extraction",
    "22": "Utilities", "23": "Construction", "31": "Manufacturing", "42": "Wholesale Trade",
    "44": "Retail Trade", "48": "Transportation and Warehousing", "51": "Information",
    "52": "Finance and Insurance", "53": "Real Estate and Rental and Leasing",
    "54": "Professional, Scientific, and Technical Services",
    "55": "Management of Companies and Enterprises",
    "56": "Administrative and Support and Waste Management", "61": "Educational Services",
    "62": "Health Care and Social Assistance", "71": "Arts, Entertainment, and Recreation",
    "72": "Accommodation and Food Services", "81": "Other Services (except Public Administration)",
    "XX": "Multi-sector companies",
}
SUBSECTOR = {  # NAICS 3-digit; (2017) / (2022) mark codes that exist in only one vintage
    "113": "Forestry and Logging", "114": "Fishing, Hunting and Trapping",
    "115": "Support Activities for Agriculture and Forestry", "211": "Oil and Gas Extraction",
    "212": "Mining (except Oil and Gas)", "213": "Support Activities for Mining",
    "221": "Utilities", "236": "Construction of Buildings",
    "237": "Heavy and Civil Engineering Construction", "238": "Specialty Trade Contractors",
    "311": "Food Manufacturing", "312": "Beverage and Tobacco Product Manufacturing",
    "313": "Textile Mills", "314": "Textile Product Mills", "315": "Apparel Manufacturing",
    "316": "Leather and Allied Product Manufacturing", "321": "Wood Product Manufacturing",
    "322": "Paper Manufacturing", "323": "Printing and Related Support Activities",
    "324": "Petroleum and Coal Products Manufacturing", "325": "Chemical Manufacturing",
    "326": "Plastics and Rubber Products Manufacturing",
    "327": "Nonmetallic Mineral Product Manufacturing", "331": "Primary Metal Manufacturing",
    "332": "Fabricated Metal Product Manufacturing", "333": "Machinery Manufacturing",
    "334": "Computer and Electronic Product Manufacturing",
    "335": "Electrical Equipment, Appliance, and Component Manufacturing",
    "336": "Transportation Equipment Manufacturing",
    "337": "Furniture and Related Product Manufacturing", "339": "Miscellaneous Manufacturing",
    "423": "Merchant Wholesalers, Durable Goods", "424": "Merchant Wholesalers, Nondurable Goods",
    "425": "Wholesale Trade Agents and Brokers", "441": "Motor Vehicle and Parts Dealers",
    "442": "Furniture and Home Furnishings Stores (2017)",
    "443": "Electronics and Appliance Stores (2017)",
    "444": "Building Material and Garden Equipment and Supplies Dealers",
    "445": "Food and Beverage Retailers", "446": "Health and Personal Care Stores (2017)",
    "447": "Gasoline Stations (2017)", "448": "Clothing and Clothing Accessories Stores (2017)",
    "449": "Furniture, Home Furnishings, Electronics, and Appliance Retailers (2022)",
    "451": "Sporting Goods, Hobby, Musical Instrument, and Book Stores (2017)",
    "452": "General Merchandise Stores (2017)", "453": "Miscellaneous Store Retailers (2017)",
    "454": "Nonstore Retailers (2017)", "455": "General Merchandise Retailers (2022)",
    "456": "Health and Personal Care Retailers (2022)",
    "457": "Gasoline Stations and Fuel Dealers (2022)",
    "458": "Clothing, Clothing Accessories, Shoe, and Jewelry Retailers (2022)",
    "459": "Sporting Goods, Hobby, Musical Instrument, Book, and Miscellaneous Retailers (2022)",
    "481": "Air Transportation", "483": "Water Transportation", "484": "Truck Transportation",
    "485": "Transit and Ground Passenger Transportation", "486": "Pipeline Transportation",
    "487": "Scenic and Sightseeing Transportation", "488": "Support Activities for Transportation",
    "492": "Couriers and Messengers", "493": "Warehousing and Storage",
    "511": "Publishing Industries (except Internet) (2017)",
    "512": "Motion Picture and Sound Recording Industries",
    "513": "Publishing Industries (2022)", "515": "Broadcasting (except Internet) (2017)",
    "516": "Broadcasting and Content Providers (2022)", "517": "Telecommunications",
    "518": "Computing Infrastructure, Data Processing, Web Hosting, and Related Services",
    "519": "Web Search Portals, Libraries, Archives, and Other Information Services",
    "522": "Credit Intermediation and Related Activities",
    "523": "Securities, Commodity Contracts, and Other Financial Investments",
    "524": "Insurance Carriers and Related Activities", "531": "Real Estate",
    "532": "Rental and Leasing Services", "533": "Lessors of Nonfinancial Intangible Assets",
    "541": "Professional, Scientific, and Technical Services",
    "551": "Management of Companies and Enterprises", "561": "Administrative and Support Services",
    "562": "Waste Management and Remediation Services", "611": "Educational Services",
    "621": "Ambulatory Health Care Services", "622": "Hospitals",
    "623": "Nursing and Residential Care Facilities", "624": "Social Assistance",
    "711": "Performing Arts, Spectator Sports, and Related Industries",
    "712": "Museums, Historical Sites, and Similar Institutions",
    "713": "Amusement, Gambling, and Recreation Industries", "721": "Accommodation",
    "722": "Food Services and Drinking Places", "811": "Repair and Maintenance",
    "812": "Personal and Laundry Services",
    "813": "Religious, Grantmaking, Civic, Professional, and Similar Organizations",
}


def log(msg):
    print(f"{dt.datetime.now():%Y-%m-%d %H:%M:%S}  {msg}", flush=True)


def fetch(name, dest):
    url = BASE + urllib.request.quote(name)
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=180) as r:
        body = r.read()
    if body[:2] != b"PK":
        raise RuntimeError(f"{url} did not return an xlsx ({len(body)} bytes)")
    dest.write_bytes(body)
    return dest


def label(level, sector=None, size=None, sub=None):
    if level == "National":
        return "United States"
    if level == "Sector":
        return SECTOR.get(sector, sector)
    if level == "Subsector":
        if sub.endswith("X"):
            return f"{SECTOR.get(sub[:2], sub[:2])} - multi-subsector companies"
        return SUBSECTOR.get(sub, sub)
    if level == "Employment size":
        return SIZE.get(size, size)
    return f"{SECTOR.get(sector, sector)} | {SIZE.get(size, size)}"


def to_num(v):
    """'23.8%' -> 0.238; 'S' / '.' / blank -> None (flag kept separately)."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (int, float)):
        return float(v) / 100 if v > 1 else float(v)
    s = str(v).strip().replace(",", "")
    if s.endswith("%"):
        try:
            return float(s[:-1]) / 100
        except ValueError:
            return None
    try:
        return float(s)
    except ValueError:
        return None


def flag(v):
    s = "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()
    return s if s in ("S", ".") else ""


def tidy(est, se, level, dims, version):
    """Census wide sheet (rows = answers, cols = smpdt) -> long rows for the AI questions."""
    est = est[pd.to_numeric(est["Question ID"], errors="coerce").isin(AI_QIDS)].copy()
    se = se[pd.to_numeric(se["Question ID"], errors="coerce").isin(AI_QIDS)].copy()
    periods = [c for c in est.columns if str(c).strip().isdigit() and len(str(c).strip()) == 6]
    key = dims + ["Question ID", "Answer"]
    for d in (est, se):
        d["Question ID"] = pd.to_numeric(d["Question ID"]).astype(int)
        for c in dims:
            d[c] = d[c].astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
        d.columns = [str(c).strip() for c in d.columns]
    periods = [str(p).strip() for p in periods]
    e = est.melt(id_vars=key, value_vars=periods, var_name="smpdt", value_name="raw_est")
    s = se.melt(id_vars=key, value_vars=periods, var_name="smpdt", value_name="raw_se")
    m = e.merge(s, on=key + ["smpdt"], how="left")
    m["flag"] = m["raw_est"].map(flag)
    m = m[m["flag"] != "."]  # '.' = question not asked in that period
    m["estimate"] = m["raw_est"].map(to_num)
    m["se"] = m["raw_se"].map(to_num)
    sec = m["Sector"] if "Sector" in m else None
    out = pd.DataFrame({
        "level": level,
        "sector": m["Sector"] if "Sector" in m else ("" if level != "Subsector" else m["Subsector"].str[:2]),
        "subsector": m["Subsector"] if "Subsector" in m else "",
        "size": m["Empsize"] if "Empsize" in m else "",
        "version": version,
        "qid": m["Question ID"],
        "measure": m["Question ID"].map(AI_QIDS),
        "answer": m["Answer"].astype(str).str.strip(),
        "smpdt": m["smpdt"].astype(int),
        "estimate": m["estimate"],
        "se": m["se"],
        "flag": m["flag"],
    })
    del sec
    out["label"] = [label(level, r.sector, r.size, r.subsector)
                    for r in out[["sector", "size", "subsector"]].itertuples()]
    return out


def read_dates(path):
    d = pd.read_excel(path, sheet_name="Collection and Reference Dates")
    d = d[pd.to_numeric(d["Smpdt"], errors="coerce").notna()]
    # current files: 'Reference Period Start' / 'Publication Date'; historical: 'Ref Start', no pub date
    col = lambda *names: next((d[n] for n in names if n in d), None)
    pub = col("Publication Date")
    return pd.DataFrame({
        "smpdt": d["Smpdt"].astype(int),
        "ref_start": pd.to_datetime(col("Reference Period Start", "Ref Start")).dt.date,
        "ref_end": pd.to_datetime(col("Ref End", "Reference Period End")).dt.date,
        "published": pd.to_datetime(pub).dt.date if pub is not None else None,
    })


def download_all(stamp_hint=None):
    """Download into raw/_new, return {name: path}. Caller renames with the publication date."""
    tmp = RAW / "_new"
    tmp.mkdir(parents=True, exist_ok=True)
    files = {}
    for level, (fname, _) in CURRENT.items():
        log(f"download {fname}")
        files[fname] = fetch(fname, tmp / fname)
    log(f"download {HIST_FILE}")
    files[HIST_FILE] = fetch(HIST_FILE, tmp / HIST_FILE)
    log(f"download {SUPP_FILE}")
    files[SUPP_FILE] = fetch(SUPP_FILE, tmp / SUPP_FILE)
    return files


def supplement(path):
    """AI Supplement 2026 (one combined estimate, 2025-11-17 .. 2026-02-08) -> long rows."""
    scopes = pd.read_excel(path, sheet_name="Data Dictionary")
    scope_lbl = {str(r[0]).strip(): str(r[1]).strip() for r in scopes.itertuples(index=False)
                 if str(r[0]).strip().isdigit()}
    rows = []
    for level, (es, ss, dim) in SUPP_SHEETS.items():
        e = pd.read_excel(path, sheet_name=es)
        s = pd.read_excel(path, sheet_name=ss)
        e = e[pd.to_numeric(e["Question ID"], errors="coerce").notna()].reset_index(drop=True)
        s = s[pd.to_numeric(s["Question ID"], errors="coerce").notna()].reset_index(drop=True)
        ecols = [c for c in e.columns if str(c).startswith("Estimate")]
        scols = [c for c in s.columns if str(c).startswith("SE")]
        for i, r in e.iterrows():
            sr = s.iloc[i] if i < len(s) else None
            for ec, sc in zip(ecols, scols):
                raw = r[ec]
                if raw is None or (isinstance(raw, float) and pd.isna(raw)):
                    continue
                code = str(r[dim]).strip().replace(".0", "") if dim else ""
                sector = code if dim == "Sector" else ""
                size = code if dim == "Empsize" else ""
                scope = str(r["Scope (see data dictionary)"]).strip().replace(".0", "")
                rows.append({
                    "level": level, "code": code,
                    "label": label(level, sector, size, ""),
                    "scope": scope_lbl.get(scope, scope),
                    "qid": int(r["Question ID"]), "question": str(r["Question"]).strip(),
                    "answer": str(r["Answer"]).strip(),
                    "column": ec.replace("Estimate - ", "").replace("Estimate", "Share"),
                    "estimate": to_num(raw), "se": to_num(sr[sc]) if sr is not None else None,
                    "flag": flag(raw),
                })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- workbook
HDR = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
V1FILL = PatternFill("solid", fgColor="EDEDED")
PCT = "0.0%"


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR
        cell.fill = HFILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def entity_key(df):
    return df["sector"].astype(str) + "|" + df["subsector"].astype(str) + "|" + df["size"].astype(str)


def wide_sheet(wb, title, df, dates, level, qid):
    """rows = entities, cols = periods (v2 newest first, then v1), Yes share."""
    d = df[(df.level == level) & (df.qid == qid) & (df.answer == "Yes")].copy()
    d["key"] = entity_key(d)
    ents = d.drop_duplicates("key")[["key", "sector", "subsector", "size", "label"]]
    ents = ents.sort_values(["sector", "subsector", "size"])
    v2p = sorted(d[d.version == V2].smpdt.unique(), reverse=True)
    v1p = sorted(d[d.version == V1].smpdt.unique(), reverse=True)
    ref = dates.set_index("smpdt")["ref_end"].to_dict()
    ws = wb.create_sheet(title)
    idcols = {"National": ["Geography"], "Sector": ["Sector", "Sector name"],
              "Subsector": ["Sector", "Subsector", "Subsector name"],
              "Employment size": ["Size code", "Employment size"],
              "Sector x Employment size": ["Sector", "Size code", "Sector | size"]}[level]
    ni = len(idcols)
    cols = [("v2", p) for p in v2p] + [("gap", None)] + [("v1", p) for p in v1p]
    ws.cell(row=1, column=1, value=f"{AI_QIDS[qid]} - share answering Yes ({level}). "
            "v2 = 'any business function' wording (Dec-2025+); grey v1 = 'producing goods or "
            "services' wording (Sep-2023..Oct-2025). Not comparable across the break. "
            "Blank = suppressed (S).").font = Font(italic=True)
    hdr1 = idcols + [("" if v == "gap" else str(p)) for v, p in cols]
    hdr2 = [""] * ni + [("" if v == "gap" else f"ref end {ref.get(p, '')}") for v, p in cols]
    hdr3 = [""] * ni + [("v2 | v1" if v == "gap" else v) for v, p in cols]
    for r, h in ((2, hdr1), (3, hdr2), (4, hdr3)):
        for c, val in enumerate(h, 1):
            ws.cell(row=r, column=c, value=val)
        style_header(ws, r, len(h))
    # fully-suppressed entities vanish from pivot_table; reindex keeps them as blank rows
    piv = d.pivot_table(index="key", columns=["version", "smpdt"], values="estimate",
                        aggfunc="first").reindex(ents.key)
    row = 5
    for e in ents.itertuples():
        ids = {"National": [e.label], "Sector": [e.sector, e.label],
               "Subsector": [e.sector, e.subsector, e.label],
               "Employment size": [e.size, e.label],
               "Sector x Employment size": [e.sector, e.size, e.label]}[level]
        for c, val in enumerate(ids, 1):
            ws.cell(row=row, column=c, value=val)
        for j, (v, p) in enumerate(cols, ni + 1):
            if v == "gap":
                continue
            vk = V2 if v == "v2" else V1
            val = piv.at[e.key, (vk, p)] if (vk, p) in piv.columns else None
            cell = ws.cell(row=row, column=j, value=None if pd.isna(val) else float(val))
            cell.number_format = PCT
            if v == "v1":
                cell.fill = V1FILL
        row += 1
    for j, (v, p) in enumerate(cols, ni + 1):
        ws.column_dimensions[get_column_letter(j)].width = 3 if v == "gap" else 9
    for c, w in zip(range(1, ni + 1), {1: [16], 2: [8, 46], 3: [8, 10, 52]}[ni]):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=5, column=ni + 1)
    return ws


def latest_sheet(wb, df, dates):
    """One row per entity: latest v2 AI use, changes, SE, expected-6m, v1 reference."""
    y = df[df.answer == "Yes"].copy()
    y["key"] = entity_key(y)
    v2p = sorted(y[y.version == V2].smpdt.unique())
    v1p = sorted(y[y.version == V1].smpdt.unique())
    last, prev = v2p[-1], v2p[-2]
    back6 = v2p[-7] if len(v2p) >= 7 else v2p[0]
    first = v2p[0]
    v1last = v1p[-1]
    ref = dates.set_index("smpdt")["ref_end"].to_dict()

    def val(level, qid, version, p):
        s = y[(y.level == level) & (y.qid == qid) & (y.version == version) & (y.smpdt == p)]
        return s.set_index("key")[["estimate", "se"]]

    ws = wb.create_sheet("Latest", 0)
    ws.cell(row=1, column=1, value=f"BTOS AI use - latest period {last} (ref period ending "
            f"{ref.get(last)}), published {dates.set_index('smpdt').published.get(last)}. "
            "Share of businesses answering Yes. pp = percentage points.").font = Font(bold=True)
    heads = ["Level", "Code", "Name",
             f"AI use {last}", "SE", f"AI use {prev}", "Δ vs prior (pp)",
             f"AI use {back6}", "Δ vs ~12 wks (pp)", f"AI use {first} (v2 start)",
             "Δ since v2 start (pp)", f"Expected next 6m {last}", "Expected − current (pp)",
             f"v1 AI use {v1last} (old wording)", "Significant vs prior? (|Δ|>1.96·√(SE²+SE²))"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=3, column=c, value=h)
    style_header(ws, 3, len(heads))
    row = 4
    for level in LEVELS:
        cur = val(level, 7, V2, last)
        pr = val(level, 7, V2, prev)
        b6 = val(level, 7, V2, back6)
        fi = val(level, 7, V2, first)
        ex = val(level, 24, V2, last)
        o1 = val(level, 7, V1, v1last)
        ents = y[(y.level == level)].drop_duplicates("key")[["key", "sector", "subsector", "size", "label"]]
        ents = ents.assign(cur=ents.key.map(cur.estimate)).sort_values("cur", ascending=False,
                                                                        na_position="last")
        for e in ents.itertuples():
            code = {"National": "US", "Sector": e.sector, "Subsector": e.subsector,
                    "Employment size": e.size,
                    "Sector x Employment size": f"{e.sector} | {e.size}"}[level]

            def g(frame, col="estimate"):
                v = frame[col].get(e.key) if e.key in frame.index else None
                return None if v is None or pd.isna(v) else float(v)
            c0, s0, p0 = g(cur), g(cur, "se"), g(pr)
            sp = g(pr, "se")
            b, f, x, o = g(b6), g(fi), g(ex), g(o1)
            dpp = lambda a, bb: None if a is None or bb is None else round((a - bb) * 100, 2)
            sig = None
            if None not in (c0, p0, s0, sp):
                sig = "yes" if abs(c0 - p0) > 1.96 * (s0 ** 2 + sp ** 2) ** 0.5 else "no"
            vals = [level, code, e.label, c0, s0, p0, dpp(c0, p0), b, dpp(c0, b), f, dpp(c0, f),
                    x, dpp(x, c0), o, sig]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=row, column=c, value=v)
                if c in (4, 5, 6, 8, 10, 12, 14):
                    cell.number_format = PCT
                elif c in (7, 9, 11, 13):
                    cell.number_format = "+0.0;-0.0;0.0"
            row += 1
    widths = [22, 12, 52, 11, 8, 11, 10, 11, 10, 12, 11, 13, 12, 13, 14]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "D4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(heads))}{row - 1}"


def national_sheet(wb, df, dates):
    """National time series (oldest -> newest) with SE + a line chart."""
    n = df[(df.level == "National") & (df.answer == "Yes")]
    ws = wb.create_sheet("National series", 1)
    heads = ["smpdt", "Ref start", "Ref end", "Published",
             "v1 AI use (goods/services)", "SE", "v1 expected 6m", "SE",
             "v2 AI use (any function)", "SE", "v2 expected 6m", "SE"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(heads))
    dd = dates.set_index("smpdt")
    piv = n.pivot_table(index="smpdt", columns=["version", "qid"], values=["estimate", "se"],
                        aggfunc="first")
    allp = sorted(set(dd.index) | set(piv.index))
    allp = [p for p in allp if p >= min(piv.index)]
    for r, p in enumerate(allp, 2):
        ws.cell(row=r, column=1, value=int(p))
        if p in dd.index:
            for c, k in ((2, "ref_start"), (3, "ref_end"), (4, "published")):
                ws.cell(row=r, column=c, value=dd.at[p, k]).number_format = "yyyy-mm-dd"
        for c, (ver, q, what) in enumerate([(V1, 7, "estimate"), (V1, 7, "se"), (V1, 24, "estimate"),
                                            (V1, 24, "se"), (V2, 7, "estimate"), (V2, 7, "se"),
                                            (V2, 24, "estimate"), (V2, 24, "se")], 5):
            v = piv[(what, ver, q)].get(p) if (what, ver, q) in piv.columns else None
            ws.cell(row=r, column=c, value=None if v is None or pd.isna(v) else float(v)).number_format = PCT
    last = len(allp) + 1
    for c, w in enumerate([8, 11, 11, 11, 12, 7, 12, 7, 12, 7, 12, 7], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "B2"
    ch = LineChart()
    ch.title = "US businesses using AI (share answering Yes) - two question wordings"
    ch.y_axis.title = "Share of businesses"
    ch.y_axis.number_format = "0%"
    ch.height, ch.width = 9, 22
    for col in (5, 7, 9, 11):
        ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=3, min_row=2, max_row=last))
    ch.x_axis.number_format = "yyyy-mm"
    ws.add_chart(ch, "N2")
    return allp


def notes_sheet(wb, dates, last, supp_ok, tracker_wave=None):
    ws = wb.create_sheet("Notes")
    pub = dates.set_index("smpdt").published.get(last)
    lines = [
        ("Census Business Trends and Outlook Survey (BTOS) - AI metrics", True),
        (f"Latest period {last}, published {pub}. Built {dt.datetime.now():%Y-%m-%d %H:%M} by btos-ai/btos_ai.py.", False),
        ("Source: https://www.census.gov/hfp/btos/data_downloads (Current Data core files + Historical "
         "'AI Core Questions (Original)' + 'AI Supplement Table 2026').", False),
        ("", False),
        ("Questions", True),
        ("Q7  'In the last two weeks, did this business use Artificial Intelligence (AI) ...?' -> 'AI use'", False),
        ("Q24 'During the next six months, do you think this business will be using AI ...?' -> 'expected 6m'", False),
        ("Answers: Yes / No / Do not know. Sheets show the Yes share; data/history.csv and the Long sheet have all answers + SEs.", False),
        ("", False),
        ("SERIES BREAK - do not splice", True),
        ("v1 (smpdt 202319-202520, Sep-2023..Oct-2025): '...use AI in producing goods or services?' (~10% Yes nationally at the end).", False),
        ("v2 (smpdt 202524 onward, Dec-2025+): '...use AI in any of its business functions?' (~18% Yes at the start).", False),
        ("The jump between them is a wording change, not adoption. Compare changes within one version only.", False),
        ("Gap: 202521-202523 (6 Oct - 16 Nov 2025) were never collected - federal funding lapse.", False),
        ("", False),
        ("Other notes", True),
        ("smpdt = ISO year + fortnight (01-26). Each estimate refers to the two-week reference period ending 'ref end'.", False),
        ("Blank cells = 'S' (suppressed by Census for quality/disclosure). Small subsectors and sector x size cells are often suppressed or noisy - check the SE.", False),
        ("Significance test on the Latest sheet: |Δ| > 1.96·√(SE₁²+SE₂²), treating the two fortnights as independent (approximate; BTOS panels overlap).", False),
        ("Subsectors use 2022 NAICS from 202616; codes 442/443/446/447/448/451-454/511/515 end then, 449/455-459/513/516 start. 519 is not comparable across 202616.", False),
        ("'XX' sector = multi-unit companies operating in several sectors; '21X', '44X' etc. = multi-subsector companies within that sector.", False),
        ("Employment size: A 1-4, B 5-9, C 10-19, D 20-49, E 50-99, F 100-249, G 250+ employees.", False),
        ("Estimates are shares of employer businesses (unweighted by employment).", False),
        ("", False),
        ("AI Supplement 2026" + ("" if supp_ok else " (not included this run - download failed)"), True),
        ("One combined estimate for 17 Nov 2025 - 8 Feb 2026 (business functions, tasks replaced, GenAI, workforce effects ...). "
         "National / sector / employment size only - Census does not publish subsector for it. 'scope' tells you the denominator.", False),
        ("", False),
        ("GenAI Adoption Tracker - worker side" + (f" (latest wave {tracker_wave:%b %Y})" if tracker_wave is not None
                                                   else " (not included this run - download failed)"), True),
        ("Source: https://www.genaiadoptiontracker.com/ (Harvard Project on Workforce; Bick, Blandin & Deming). "
         "Real-Time Population Survey of US adults 18-64, quarterly waves (Aug / Nov / Feb / May), ~5,000 per wave.", False),
        ("Pulled from the CSVs the site's charts load (GenAI_All / _Ind / _Occ / _HoursShare_* / _TimeSav_* / _PC_Internet).", False),
        ("Industry / Occupation = share of EMPLOYED respondents using genAI for work. Hours share and time saved are % of all "
         "work hours (users and non-users). The Aug 2024 wave did not ask hours / time saved (blank).", False),
        ("BTOS counts firms (any AI, last 2 weeks); the tracker counts people (genAI, ever/last week). Levels differ by design.", False),
        ("", False),
        ("Update: Windows task 'AI Adoption Monitor Weekly' (Fridays 07:00) runs btos-ai/run_weekly_task.vbs; "
         "a new workbook appears only when Census publishes a new fortnight or the tracker CSVs change.", False),
    ]
    for r, (t, b) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=t).font = Font(bold=b, size=12 if r == 1 else 11)
    ws.column_dimensions["A"].width = 150


def write_frame(ws, frame, pct_cols=(), start_row=1, widths=None):
    for c, h in enumerate(frame.columns, 1):
        ws.cell(row=start_row, column=c, value=str(h))
    style_header(ws, start_row, len(frame.columns))
    for r, rec in enumerate(frame.itertuples(index=False), start_row + 1):
        for c, v in enumerate(rec, 1):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            v = None if (v is None or (isinstance(v, float) and pd.isna(v))) else v
            cell = ws.cell(row=r, column=c, value=v)
            if frame.columns[c - 1] in pct_cols:
                cell.number_format = PCT
            elif isinstance(v, dt.datetime):
                cell.number_format = "mmm yyyy"
    for c, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return start_row + len(frame) + 1


def tracker_sheets(wb, gt, pcnet):
    """GenAI Adoption Tracker (worker survey) sheets."""
    nat = gt[gt.dimension == "National"].copy()
    nat["series"] = nat.category + " | " + nat.measure
    order = list(dict.fromkeys(nat.series))
    piv = nat.pivot_table(index="wave", columns="series", values=["mean", "se"], aggfunc="first")
    rows = pd.DataFrame({"Survey wave": piv.index})
    for s in order:
        rows[s] = piv[("mean", s)].values
        rows[s + " SE"] = piv[("se", s)].values
    ws = wb.create_sheet("Workers GenAI - National")
    ws.cell(row=1, column=1, value="GenAI Adoption Tracker (Bick, Blandin & Deming, Real-Time "
            "Population Survey): US adults 18-64, quarterly. Shares of PEOPLE, not firms.").font = Font(bold=True)
    end = write_frame(ws, rows, pct_cols=set(rows.columns[1:]), start_row=2,
                      widths=[11] + [13] * (len(rows.columns) - 1))
    for c in range(2, len(rows.columns) + 1):
        ws.cell(row=2, column=c).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 75
    ws.freeze_panes = "B3"
    ch = LineChart()
    ch.title = "Share of US adults / workers using genAI"
    ch.y_axis.number_format = "0%"
    ch.height, ch.width = 9, 20
    for s in ("All adults - Any | Share Using GenAI", "Employed - For Work | Share Using GenAI",
              "All adults - Non-Work | Share Using GenAI"):
        if s in rows.columns:
            col = list(rows.columns).index(s) + 1
            ch.add_data(Reference(ws, min_col=col, min_row=2, max_row=end - 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=3, max_row=end - 1))
    ch.x_axis.number_format = "mmm yy"
    ws.add_chart(ch, f"A{end + 2}")

    for dim in ("Industry", "Occupation"):
        ws = wb.create_sheet(f"Workers GenAI - {dim}")
        d = gt[gt.dimension == dim]
        row = 1
        for meas in ("Share of workers using genAI at work", "Share of work hours using genAI",
                     "Time saved by genAI (% of work hours)"):
            m = d[d.measure == meas]
            if m.empty:
                continue
            p = m.pivot_table(index="category", columns="wave", values="mean", aggfunc="first")
            p = p[sorted(p.columns, reverse=True)]
            se_last = m[m.wave == m.wave.max()].set_index("category").se
            p = p.sort_values(p.columns[0], ascending=False)
            frame = pd.DataFrame({dim: p.index})
            if dim == "Industry":
                frame.insert(0, "NAICS", [gtmod.naics(c) for c in p.index])
            for w in p.columns:
                frame[w.strftime("%b %Y")] = p[w].values
            frame[f"SE {p.columns[0]:%b %Y}"] = se_last.reindex(p.index).values
            ws.cell(row=row, column=1, value=meas).font = Font(bold=True, size=12)
            ni = 2 if dim == "Industry" else 1
            row = write_frame(ws, frame, pct_cols=set(frame.columns[ni:]), start_row=row + 1,
                              widths=([7, 52] if ni == 2 else [58]) + [10] * (len(frame.columns) - ni))
            row += 2
        ws.freeze_panes = "C2" if dim == "Industry" else "B2"

    if pcnet is not None:
        ws = wb.create_sheet("Workers GenAI - PC & Internet")
        ws.cell(row=1, column=1, value="Adoption by years since mass-market launch: genAI (RPS) vs "
                "personal computer (CPS) vs internet (CPS / ITU). Share of adults.").font = Font(bold=True)
        write_frame(ws, pcnet, pct_cols=set(pcnet.columns[1:]), start_row=2,
                    widths=[14] + [13] * (len(pcnet.columns) - 1))


def firm_vs_worker_sheet(wb, df, gt, dates):
    """BTOS firm AI use (latest v2 fortnight) vs tracker worker genAI use (latest wave), by sector."""
    y = df[(df.level == "Sector") & (df.answer == "Yes") & (df.version == V2)]
    last = int(y.smpdt.max())
    cur = y[(y.smpdt == last) & (y.qid == 7)].set_index("sector")
    exp = y[(y.smpdt == last) & (y.qid == 24)].set_index("sector")
    natb = df[(df.level == "National") & (df.answer == "Yes") & (df.version == V2) &
              (df.smpdt == last)].set_index("qid").estimate
    ind = gt[gt.dimension == "Industry"]
    wl = ind.wave.max()
    t = ind[ind.wave == wl].pivot_table(index="naics", columns="measure", values="mean", aggfunc="first")
    natw = gt[(gt.dimension == "National") & (gt.wave == wl)]
    work_us = natw[(natw.category == "Employed - For Work") & (natw.measure == "Share Using GenAI")]["mean"]
    hrs_us = natw[natw.file == "GenAI_HoursShare_All"]["mean"]
    tsv_us = natw[natw.file == "GenAI_TimeSav_All"]["mean"]
    codes = [c for c in SECTOR if c != "XX"] + ["92"]
    rows = [{"NAICS": "US", "Sector": "All sectors",
             "Firms using AI (BTOS)": natb.get(7), "Firms expecting AI in 6m (BTOS)": natb.get(24),
             "Workers using genAI at work (tracker)": work_us.iloc[0] if len(work_us) else None,
             "Work hours using genAI (tracker)": hrs_us.iloc[0] if len(hrs_us) else None,
             "Time saved, % of hours (tracker)": tsv_us.iloc[0] if len(tsv_us) else None}]
    for c in codes:
        g = lambda frame: (frame.estimate.get(c) if c in frame.index else None)
        tv = lambda meas: (t.at[c, meas] if c in t.index and meas in t.columns else None)
        rows.append({"NAICS": c, "Sector": SECTOR.get(c, "Public Administration"),
                     "Firms using AI (BTOS)": g(cur), "Firms expecting AI in 6m (BTOS)": g(exp),
                     "Workers using genAI at work (tracker)": tv("Share of workers using genAI at work"),
                     "Work hours using genAI (tracker)": tv("Share of work hours using genAI"),
                     "Time saved, % of hours (tracker)": tv("Time saved by genAI (% of work hours)")})
    f = pd.DataFrame(rows)
    f["Worker ÷ firm share"] = [None if (a is None or b is None or pd.isna(a) or pd.isna(b) or not b)
                                else round(a / b, 2) for a, b in
                                zip(f["Workers using genAI at work (tracker)"], f["Firms using AI (BTOS)"])]
    f = pd.concat([f.iloc[:1], f.iloc[1:].sort_values("Firms using AI (BTOS)", ascending=False,
                                                     na_position="last")])
    ws = wb.create_sheet("Firm vs worker (sector)", 2)
    ref = dates.set_index("smpdt").ref_end.get(last)
    ws.cell(row=1, column=1, value=f"Firm vs worker AI adoption by sector - BTOS {last} (fortnight "
            f"ending {ref}) vs GenAI Adoption Tracker {wl:%b %Y} wave").font = Font(bold=True)
    ws.cell(row=2, column=1, value="Not the same question: BTOS = share of FIRMS that used any AI in a "
            "business function in the last 2 weeks (each firm counts once). Tracker = share of WORKERS "
            "who use genAI for work. Rank agreement is the useful signal, not the level gap.").font = Font(italic=True)
    end = write_frame(ws, f, pct_cols={c for c in f.columns[2:7]}, start_row=4,
                      widths=[7, 46, 12, 13, 14, 13, 13, 10])
    ws.row_dimensions[4].height = 45
    for c in range(1, len(f.columns) + 1):
        ws.cell(row=4, column=c).alignment = Alignment(wrap_text=True, vertical="center")
    both = f.iloc[1:].dropna(subset=["Firms using AI (BTOS)", "Workers using genAI at work (tracker)"])
    if len(both) > 4:
        rho = both["Firms using AI (BTOS)"].rank().corr(both["Workers using genAI at work (tracker)"].rank())
        ws.cell(row=end + 1, column=1, value=f"Spearman rank correlation across {len(both)} sectors: "
                f"{rho:.2f}").font = Font(bold=True)
    ws.freeze_panes = "C5"


FED_URL = ("https://www.federalreserve.gov/econres/notes/feds-notes/"
           "monitoring-ai-adoption-in-the-u-s-economy-20260403.html")
FED_INDUSTRIES = [("54", "Professional services"), ("52", "Finance and insurance"),
                  ("62", "Health care and social assistance"), ("31", "Manufacturing"),
                  ("44", "Retail trade")]
SBU_CSV = HERE / "static" / "sbu_2025-11_ai_technology.csv"


def fed_series(df):
    """The Fed note's BTOS cuts, every period: 4-observation moving average within each wording
    version (footnote 2), size 1-49 = simple average of classes A-D (figure 3 note)."""
    y = df[df.answer == "Yes"]
    cols = {}

    def add(name, frame):
        cols[name] = frame.groupby(["version", "smpdt"]).estimate.mean()
    nat = y[y.level == "National"]
    add("US: AI use", nat[nat.qid == 7])
    add("US: expected 6m", nat[nat.qid == 24])
    sz = y[(y.level == "Employment size") & (y.qid == 7)]
    add("Size 1-49", sz[sz["size"].isin(list("ABCD"))])
    for code, lbl in (("E", "Size 50-99"), ("F", "Size 100-249"), ("G", "Size 250+")):
        add(lbl, sz[sz["size"] == code])
    sec = y[(y.level == "Sector") & (y.qid == 7)]
    for code, lbl in FED_INDUSTRIES:
        add(lbl, sec[sec.sector == code])
    raw = pd.DataFrame(cols).sort_index()
    ma = raw.groupby(level="version", group_keys=False).apply(lambda g: g.rolling(4, min_periods=4).mean())
    return raw, ma


def fed_tables(df, gt, dates):
    """Allen (2026) FEDS Note tables rebuilt on live data. Returns {name: DataFrame} + 'meta'.
    Shared by the workbook (fed_note_sheet) and the web app."""
    raw, ma = fed_series(df)
    v2 = ma.loc[V2]
    v1 = ma.loc[V1].dropna(how="all")
    # Fed's year-end-2025 point = MA of the first 4 v2 periods (index BEFORE dropping the NaN warm-up)
    last, first4 = v2.index.max(), v2.index[3]
    ref = dates.set_index("smpdt").ref_end.to_dict()
    sbu = pd.read_csv(SBU_CSV) if SBU_CSV.exists() else None

    def sbu_val(cohort_start, tech):
        if sbu is None:
            return None
        s = sbu[sbu.cohort.str.startswith(cohort_start) & (sbu.technology == tech)]
        return float(s.currently_using.iloc[0]) if len(s) else None

    nw = gt[gt.dimension == "National"] if gt is not None else pd.DataFrame()

    def rps(cat, meas, wave=None):
        if nw.empty:
            return None, None
        s = nw[(nw.category == cat) & (nw.measure == meas)].set_index("wave")["mean"].sort_index()
        if wave is None:
            wave = s.index.max()
        return (float(s.get(wave)) if wave in s.index else None), wave

    out = {}
    w_any, wl = rps("Employed - For Work", "Share Using GenAI")
    out["t1"] = pd.DataFrame([
        ["AI", "BTOS", f"MA4 to {last} (ref end {ref.get(last)})", "Firms", "Firm-weighted", v2.at[last, "US: AI use"], 0.18],
        ["AI", "SBU", "Nov 2025 (one-off)", "Firms", "Employment-weighted", sbu_val("Full sample (Employment", "Any AI"), 0.78],
        ["GenAI/LLMs", "RPS", f"{wl:%b %Y}" if wl is not None else "", "Individuals", "Share of labor force", w_any, 0.41],
        ["GenAI/LLMs", "SBU", "Nov 2025 (one-off)", "Firms", "Employment-weighted", sbu_val("Full sample (Employment", "Text generation (LLM)"), 0.54],
    ], columns=["Scope", "Source", "Date", "Unit", "Estimation", "Latest", "Fed note (Dec/Nov 2025)"])

    nov25 = pd.Timestamp("2025-11-01")
    chk = [
        ["BTOS firms using AI, new wording, MA4 of first 4 periods (year-end 2025)", v2.at[first4, "US: AI use"], 0.18],
        ["BTOS firms expecting AI in 6m, same", v2.at[first4, "US: expected 6m"], 0.21],
        ["BTOS firms using AI, old wording, last MA4", v1["US: AI use"].iloc[-1], 0.10],
        ["BTOS firms expecting AI, old wording, last MA4", v1["US: expected 6m"].iloc[-1], 0.14],
        ["BTOS professional services, new wording, year-end 2025", v2.at[first4, "Professional services"], 0.33],
        ["BTOS finance and insurance, new wording, year-end 2025", v2.at[first4, "Finance and insurance"], 0.30],
        ["RPS workers using genAI at work, Nov 2025", rps("Employed - For Work", "Share Using GenAI", nov25)[0], 0.407],
        ["RPS used at work at least once last week, Nov 2025", rps("Employed - For Work", "Share Used GenAI Last Week", nov25)[0], 0.352],
        ["RPS used at work every day last week, Nov 2025", rps("Employed - For Work", "Share Used GenAI Every Day Last Week", nov25)[0], 0.12],
        ["RPS non-work genAI use, Nov 2025", rps("All adults - Non-Work", "Share Using GenAI", nov25)[0], 0.50],
    ]
    t = pd.DataFrame(chk, columns=["Metric", "This workbook", "Fed note"])
    t["Diff (pp)"] = [None if a is None or pd.isna(a) else round((a - b) * 100, 1)
                      for a, b in zip(t["This workbook"], t["Fed note"])]
    out["check"] = t

    if not nw.empty:
        yago = wl - pd.DateOffset(years=1)
        t2 = []
        for lbl, meas in (("Any", "Share Using GenAI"), ("Daily last week", "Share Used GenAI Every Day Last Week"),
                          ("At least once last week", "Share Used GenAI Last Week")):
            a = rps("Employed - For Work", meas, wl)[0]
            b = rps("Employed - For Work", meas, yago)[0]
            t2.append([lbl, a, None if None in (a, b) else round((a - b) * 100, 1),
                       None if None in (a, b) or not b else a / b - 1])
        out["t2"] = pd.DataFrame(t2, columns=["Frequency", "Level", "YoY chg (pp)", "YoY % chg"])

    sizes = [("Size 1-49", "0-50"), ("Size 50-99", "50-99"), ("Size 100-249", "99-249"), ("Size 250+", "250+")]
    out["t3"] = pd.DataFrame([[s.replace("Size ", ""), v1[s].iloc[-1], v2.at[first4, s], v2.at[last, s],
                               sbu_val(c, "Any AI"), sbu_val(c, "Text generation (LLM)")] for s, c in sizes],
                             columns=["Employees", "BTOS old wording (last MA4)", "BTOS new, year-end 2025",
                                      f"BTOS new, latest ({last})", "SBU any AI", "SBU LLM"])

    ind = gt[gt.dimension == "Industry"] if gt is not None else pd.DataFrame()

    def rps_ind(code, wave):
        if ind.empty or wave is None:
            return None
        s = ind[(ind.naics == code) & (ind.measure == "Share of workers using genAI at work") & (ind.wave == wave)]["mean"]
        return float(s.iloc[0]) if len(s) else None
    t4 = []
    for code, lbl in FED_INDUSTRIES:
        a = rps_ind(code, wl)
        b = rps_ind(code, wl - pd.DateOffset(years=1)) if wl is not None else None
        t4.append([lbl, v1[lbl].iloc[-1], v2.at[first4, lbl], v2.at[last, lbl], a,
                   None if None in (a, b) else round((a - b) * 100, 1)])
    out["t4"] = pd.DataFrame(t4, columns=["Industry", "BTOS old wording (last MA4)", "BTOS new, year-end 2025",
                                          f"BTOS new, latest ({last})",
                                          f"RPS at work {wl:%b %Y}" if wl is not None else "RPS", "RPS YoY (pp)"])

    ts = ma.reset_index()
    ts.insert(2, "ref end", ts.smpdt.map(ref))
    ts["version"] = ts.version.map({V1: "v1 old", V2: "v2 new"})
    out["ts"] = ts.sort_values("smpdt").reset_index(drop=True)
    out["sbu"] = sbu
    out["meta"] = {"last": int(last), "first4": int(first4), "wave": wl}
    return out


def fed_note_sheet(wb, df, gt, dates):
    """Allen (2026) FEDS Note 'Monitoring AI Adoption in the US Economy' rebuilt on live data."""
    T = fed_tables(df, gt, dates)
    m = T["meta"]
    ws = wb.create_sheet("Fed note (Allen 2026)", 3)
    ws.cell(row=1, column=1, value="Fed FEDS Note, 3 Apr 2026 - 'Monitoring AI Adoption in the US "
            "Economy' (J. Allen): its tables and figures rebuilt on the latest BTOS + RPS data").font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value=FED_URL)
    ws.cell(row=3, column=1, value="Fed method: BTOS = 4-period moving average (MA4) within one question wording; "
            "size 1-49 = simple average of the 1-4/5-9/10-19/20-49 classes. RPS = GenAI Adoption Tracker. "
            "SBU = Atlanta Fed Survey of Business Uncertainty, asked once (Nov 2025), employment-weighted.").font = Font(italic=True)
    wl = m["wave"]
    blocks = [
        ("t1", "Table 1 - summary of estimated AI adoption rates (live where possible)", 5),
        ("check", "Replication check - this workbook's data at the Fed note's dates vs the numbers printed in the note", 1),
        ("t2", f"Table 2 - genAI use at work, intensity (RPS, {wl:%b %Y} vs a year earlier)" if wl is not None else "", None),
        ("t3", "Figure 3 - AI adoption by firm size: BTOS MA4 (old wording last / new wording first-4 / latest) vs SBU Nov 2025", 1),
        ("t4", "Figure 4 - five peer industries: BTOS firm AI use (MA4) and RPS worker genAI use at work", 1),
        ("ts", "Figures 2-4, BTOS series: MA4 by period (blank until 4 observations exist within a wording)", 3),
    ]
    row = 5
    for key, title, first_pct in blocks:
        if key not in T:
            continue
        f = T[key]
        ws.cell(row=row, column=1, value=title).font = Font(bold=True)
        if key == "t2":
            pct = {"Level", "YoY % chg"}
        elif key == "t4":
            pct = set(f.columns[1:5])
        elif key == "check":
            pct = {"This workbook", "Fed note"}
        elif key == "t1":
            pct = {"Latest", "Fed note (Dec/Nov 2025)"}
        else:
            pct = set(f.columns[first_pct:])
        row = write_frame(ws, f, pct_cols=pct, start_row=row + 1) + 1
    for c, w in enumerate([62, 14, 20, 12, 18, 14, 14, 13, 13, 13, 13, 13, 13], 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    sbu = T["sbu"]
    if sbu is not None:
        ws2 = wb.create_sheet("SBU Nov 2025 (static)")
        ws2.cell(row=1, column=1, value="Atlanta Fed Survey of Business Uncertainty, Nov 10-21 2025: firm AI technology "
                 "use, % of firms, employment-weighted (except the equal-weighted row). Source: Dec 2025 SBU monthly "
                 "report, slide 9. One-off - not updated.").font = Font(bold=True)
        p = sbu.pivot_table(index=["group", "cohort", "n"], columns="technology", values="currently_using",
                            sort=False).reset_index()
        write_frame(ws2, p, pct_cols=set(p.columns[3:]), start_row=3,
                    widths=[12, 44, 6] + [13] * (len(p.columns) - 3))
        p2 = sbu.pivot_table(index=["group", "cohort", "n"], columns="technology", values="expect_next_3y",
                             sort=False).reset_index()
        r2 = 3 + len(p) + 3
        ws2.cell(row=r2 - 1, column=1, value="Expect to use over the next 3 years").font = Font(bold=True)
        write_frame(ws2, p2, pct_cols=set(p2.columns[3:]), start_row=r2)


def build_workbook(df, dates, supp, path, gt=None, pcnet=None):
    wb = Workbook()
    wb.remove(wb.active)
    last = int(df[df.version == V2].smpdt.max())
    latest_sheet(wb, df, dates)
    national_sheet(wb, df, dates)
    if gt is not None and len(gt):
        firm_vs_worker_sheet(wb, df, gt, dates)
        tracker_sheets(wb, gt, pcnet)
    try:
        fed_note_sheet(wb, df, gt, dates)
    except Exception as ex:  # a derived view; never block the workbook
        log(f"Fed note sheet skipped: {ex!r}")
    short = {"Sector": "Sector", "Subsector": "Subsector", "Employment size": "Size",
             "Sector x Employment size": "Sector x Size"}
    for level in LEVELS[1:]:
        wide_sheet(wb, f"{short[level]} - AI use", df, dates, level, 7)
        wide_sheet(wb, f"{short[level]} - Expected 6m", df, dates, level, 24)
    if supp is not None and len(supp):
        ws = wb.create_sheet("AI Supplement 2026")
        cols = ["level", "code", "label", "scope", "qid", "question", "answer", "column",
                "estimate", "se", "flag"]
        for c, h in enumerate(cols, 1):
            ws.cell(row=1, column=c, value=h)
        style_header(ws, 1, len(cols))
        for r, rec in enumerate(supp[cols].itertuples(index=False), 2):
            for c, v in enumerate(rec, 1):
                v = None if (isinstance(v, float) and pd.isna(v)) else v
                cell = ws.cell(row=r, column=c, value=v)
                if c in (9, 10):
                    cell.number_format = PCT
        for c, w in enumerate([15, 6, 40, 40, 5, 70, 50, 12, 9, 8, 5], 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:K{len(supp) + 1}"
    long = wb.create_sheet("Long (all answers)")
    lc = ["level", "sector", "subsector", "size", "label", "version", "measure", "answer", "smpdt",
          "ref_end", "published", "estimate", "se", "flag"]
    lf = df.merge(dates[["smpdt", "ref_end", "published"]], on="smpdt", how="left")
    lf = lf[lc].sort_values(["level", "sector", "subsector", "size", "measure", "answer", "smpdt"])
    for c, h in enumerate(lc, 1):
        long.cell(row=1, column=c, value=h)
    style_header(long, 1, len(lc))
    for r, rec in enumerate(lf.itertuples(index=False), 2):
        for c, v in enumerate(rec, 1):
            v = None if (v is None or (isinstance(v, float) and pd.isna(v))) else v
            cell = long.cell(row=r, column=c, value=v)
            if c in (12, 13):
                cell.number_format = PCT
            elif c in (10, 11):
                cell.number_format = "yyyy-mm-dd"
    long.freeze_panes = "A2"
    long.auto_filter.ref = f"A1:N{len(lf) + 1}"
    notes_sheet(wb, dates, last, supp is not None and len(supp) > 0,
                gt.wave.max() if gt is not None and len(gt) else None)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="rebuild even if no new period")
    ap.add_argument("--keep", type=int, default=8, help="dated workbooks to keep in output/")
    a = ap.parse_args()
    for p in (RAW, DATA, OUT):
        p.mkdir(parents=True, exist_ok=True)
    state = json.loads(STATE.read_text()) if STATE.exists() else {}

    # cheap check first: the 90 KB National file tells us the latest period
    tmp = RAW / "_new"
    tmp.mkdir(exist_ok=True)
    nat = fetch("National.xlsx", tmp / "National.xlsx")
    # the dates sheet lists Census's future schedule too, so take the newest estimate column
    hdr = pd.read_excel(nat, sheet_name="Response Estimates", nrows=0).columns
    latest = max(int(str(c).strip()) for c in hdr if str(c).strip().isdigit() and len(str(c).strip()) == 6)
    try:
        traw = gtmod.fetch_all()
        tdig = gtmod.digest(traw)
    except Exception as ex:  # tracker is a second source; never block the BTOS build
        log(f"GenAI tracker download failed: {ex!r}")
        traw, tdig = None, state.get("tracker_digest")
    out_exists = state.get("workbook") and (HERE / state["workbook"]).exists()
    if (not a.force and state.get("latest_smpdt") == latest and out_exists
            and state.get("tracker_digest") == tdig):
        log(f"no new data (BTOS {latest}, tracker {tdig} already processed) - nothing to do")
        return 0

    files = download_all()
    cur_dates = read_dates(files["National.xlsx"])
    hist_dates = read_dates(files[HIST_FILE])
    dates = (pd.concat([hist_dates, cur_dates]).drop_duplicates("smpdt", keep="last")
             .sort_values("smpdt").reset_index(drop=True))

    parts = []
    for level, (fname, dims) in CURRENT.items():
        est = pd.read_excel(files[fname], sheet_name="Response Estimates", dtype=object)
        se = pd.read_excel(files[fname], sheet_name="Response Standard Errors", dtype=object)
        parts.append(tidy(est, se, level, dims, V2))
        log(f"{level}: {len(parts[-1])} v2 rows")
    for level, (es, ss, dims) in HIST_SHEETS.items():
        est = pd.read_excel(files[HIST_FILE], sheet_name=es, dtype=object)
        se = pd.read_excel(files[HIST_FILE], sheet_name=ss, dtype=object)
        parts.append(tidy(est, se, level, dims, V1))
        log(f"{level}: {len(parts[-1])} v1 rows")
    df = pd.concat(parts, ignore_index=True)
    # Census re-sends the whole history every release; keep any period it later drops
    hist_csv = DATA / "history.csv"
    keycols = ["level", "sector", "subsector", "size", "version", "qid", "answer", "smpdt"]
    if hist_csv.exists():
        old = pd.read_csv(hist_csv, dtype={"sector": str, "subsector": str, "size": str})
        old = old.drop(columns=[c for c in ("ref_start", "ref_end", "published") if c in old])
        for c in ("sector", "subsector", "size"):
            old[c] = old[c].fillna("")
        old = old[df.columns]
        merged = pd.concat([df, old], ignore_index=True)
        df = merged.drop_duplicates(keycols, keep="first").reset_index(drop=True)
    for c in ("sector", "subsector", "size"):
        df[c] = df[c].fillna("").astype(str)

    try:
        supp = supplement(files[SUPP_FILE])
    except Exception as ex:  # supplement is a bonus; never block the core build
        log(f"AI supplement skipped: {ex!r}")
        supp = None

    gt = pcnet = None
    gcsv = DATA / "genai_tracker_history.csv"
    if traw is not None:
        try:
            gt = gtmod.tidy(traw)
            pcnet = gtmod.pc_internet(traw)
            gk = ["file", "dimension", "category", "measure", "wave"]
            if gcsv.exists():  # keep any wave the site later drops
                old = pd.read_csv(gcsv, parse_dates=["wave"], dtype={"naics": str}).fillna({"naics": ""})
                gt = pd.concat([gt, old[gt.columns]]).drop_duplicates(gk, keep="first")
            gt = gt.sort_values(gk).reset_index(drop=True)
            gt.to_csv(gcsv, index=False)
            pcnet.to_csv(DATA / "genai_pc_internet.csv", index=False)
            log(f"GenAI tracker: {len(gt)} rows, latest wave {gt.wave.max():%b %Y}")
        except Exception as ex:
            log(f"GenAI tracker parse failed: {ex!r}")
            gt = pcnet = None
    if gt is None and gcsv.exists():
        gt = pd.read_csv(gcsv, parse_dates=["wave"], dtype={"naics": str}).fillna({"naics": ""})

    last = int(df[df.version == V2].smpdt.max())
    pub = dates.set_index("smpdt").published.get(last)
    stamp = pd.Timestamp(pub).strftime("%Y%m%d")
    out = OUT / f"{stamp}_{OUTNAME}.xlsx"
    build_workbook(df, dates, supp, out, gt, pcnet)
    df.merge(dates, on="smpdt", how="left").sort_values(keycols).to_csv(hist_csv, index=False)

    # raw: keep only the latest publication, named YYYYMMDD_<file>
    for f in RAW.glob("*.xlsx"):
        f.unlink()
    for name, p in files.items():
        os.replace(p, RAW / f"{stamp}_{name}")
    for p in tmp.glob("*"):
        p.unlink()
    olds = sorted(OUT.glob(f"*_{OUTNAME}.xlsx"))
    for f in olds[:-a.keep]:
        f.unlink()

    nat_yes = df[(df.level == "National") & (df.qid == 7) & (df.answer == "Yes") &
                 (df.version == V2)].sort_values("smpdt")
    tail = ", ".join(f"{int(r.smpdt)} {r.estimate:.1%}" for r in nat_yes.tail(3).itertuples())
    log(f"wrote {out.name}  (latest {last}, published {pub}); US AI use v2: {tail}")
    STATE.write_text(json.dumps({"latest_smpdt": last, "published": str(pub), "tracker_digest": tdig,
                                 "workbook": str(out.relative_to(HERE)),
                                 "built": dt.datetime.now().isoformat(timespec="seconds")}, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
